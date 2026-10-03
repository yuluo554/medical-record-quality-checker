"""报告导出层（docx）。依赖属可选 extra，惰性导入：缺依赖给安装提示不崩。

报告结构（plan/05 M5；参考 eval/detect.render_markdown 的报告结构）：
首页免责声明（非临床决策支持 + 合成数据声明）→ 患者基本信息 → 质控汇总
→ 结论清单（分级 + 依据关联：文号+条款+摘录）→ 通过项简表 → 签署栏。
0 外链纪律：docx 内不得引用外部资源（不加超链接/外链图片），守门测试扫描
docx 包内 rels 断言无 External 关系。
"""

from pathlib import Path
from typing import List, Optional

from ..pipeline.stages import STATUS_LABEL, summarize

__all__ = ["export_report", "docx_available"]

_DISCLAIMER = ("免责声明：本报告由病历内涵质控智能审核系统自动生成，只做病历书写质量与"
               "记录一致性质控，不是临床决策支持工具；评测与演示数据全部为合成数据，"
               "结论需经质控医师人工复核签字后方可使用，不构成诊疗依据。")


def _import_docx():
    """惰性导入 python-docx；缺失时抛带安装提示的 ImportError。"""
    try:
        import docx  # noqa: F401
    except ImportError as exc:
        raise ImportError(
            "质控报告导出需要 python-docx：请安装报告组件 pip install 'mrqc[report]'"
        ) from exc


def docx_available() -> bool:
    """docx 依赖是否可用（CLI 预检用，避免跑完整质控后才在导出节点失败）。"""
    try:
        import docx  # noqa: F401
        return True
    except ImportError:
        return False


def _fv_text(fv) -> str:
    """FieldValue → 展示文本（值 + 单位；空值显示 —）。"""
    if fv is None or fv.value is None or fv.value == "":
        return "—"
    text = str(fv.value)
    if fv.unit and not text.endswith(fv.unit):
        text += fv.unit
    return text


def _set_cn_font(doc) -> None:
    """正文中文用宋体（标题走模板默认），字体美化失败不阻塞导出。"""
    try:
        from docx.shared import Pt
        from docx.oxml.ns import qn

        normal = doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(10.5)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    except Exception:  # noqa: BLE001 —— 字体美化属锦上添花
        pass


def _add_basis_line(par, finding) -> None:
    """依据关联行：文号 + 条款 + 摘录 + 核对状态（0 外链，纯文本）。"""
    basis = finding.basis
    if basis is None:
        run = par.add_run("依据：未关联")
        run.italic = True
        return
    par.add_run("依据：%s %s %s（%s）" % (basis.document, basis.document_no,
                                         basis.clause, basis.status))
    if basis.quote:
        quote = basis.quote.replace("\n", " ")
        if len(quote) > 100:
            quote = quote[:100] + "…"
        lead = par.add_run("｜条文摘录：")
        lead.bold = True
        par.add_run("「%s」" % quote)


def export_report(card, findings: List, out_path, summary: Optional[dict] = None) -> Path:
    """导出 docx 质控报告：首页免责声明 + 基本信息 + 分级统计 + 依据关联结论清单 + 签署栏。

    summary 缺省时由 findings 现算（n_rules 未知记 —）；流水线节点调用时传入
    stages.summarize 的结果保证与 CLI/Web 输出同源。
    """
    _import_docx()
    from docx import Document
    from docx.shared import Pt

    summary = summary if summary is not None else summarize(findings)
    out_path = Path(out_path)
    if str(out_path.parent) not in ("", "."):
        out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    _set_cn_font(doc)

    doc.add_heading("病历内涵质控报告", level=0)
    dis = doc.add_paragraph()
    run = dis.add_run(_DISCLAIMER)
    run.bold = True
    run.font.size = Pt(9)

    # 一、患者基本信息（全部来自参数卡；合成数据亦如实展示）
    doc.add_heading("一、患者基本信息", level=1)
    p = card.patient
    basic_rows = [
        ("姓名", _fv_text(p.name)), ("性别", _fv_text(p.gender)),
        ("年龄", _fv_text(p.age)), ("科室", _fv_text(p.department)),
        ("床号", _fv_text(p.bed)), ("入院时间", _fv_text(p.admission_date)),
        ("出院时间", _fv_text(p.discharge_date)), ("住院天数", _fv_text(p.hospital_days)),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in basic_rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    # 二、质控汇总（分级统计）
    doc.add_heading("二、质控汇总", level=1)
    sev = summary.get("non_pass_severity") or {}
    sev_text = "；".join("%s %d 条" % (k, v) for k, v in sev.items()) or "无"
    n_rules = summary.get("n_rules")
    stat_rows = [
        ("参评规则", "%s 条" % (n_rules if n_rules is not None else "—")),
        ("结论总数", "%d 条（通过 %d / 不通过 %d / 待人工确认 %d）"
         % (summary["n_findings"], summary["n_pass"], summary["n_fail"],
            summary["n_need_confirm"])),
        ("总体结论", summary.get("overall_label", "—")),
        ("非 pass 结论分级分布", sev_text),
        ("依据关联率", "%.1f%%" % (summary.get("basis_linkage_rate", 0.0) * 100)),
        ("解析告警", "%d 条" % len(card.parse_warnings)),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in stat_rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    # 三、结论清单（非 pass，逐条含分级/依据关联/病历证据/建议）
    doc.add_heading("三、结论清单（非通过项）", level=1)
    non_pass = [f for f in findings if f.status != "pass"]
    if not non_pass:
        doc.add_paragraph("全部规则通过，无非通过结论。")
    for idx, f in enumerate(non_pass, 1):
        head = doc.add_paragraph()
        head.add_run("%d. %s %s" % (idx, f.rule_id, f.rule_name)).bold = True
        head.add_run("　［%s｜%s］" % (f.severity.label,
                                      STATUS_LABEL.get(f.status, f.status)))
        line = doc.add_paragraph()
        line.paragraph_format.left_indent = Pt(18)
        _add_basis_line(line, f)
        for ev in f.evidence:
            line = doc.add_paragraph()
            line.paragraph_format.left_indent = Pt(18)
            lead = line.add_run("病历证据：")
            lead.bold = True
            line.add_run("[%s] %s" % (ev.part, ev.quote))
        if f.suggestion:
            line = doc.add_paragraph()
            line.paragraph_format.left_indent = Pt(18)
            lead = line.add_run("整改建议：")
            lead.bold = True
            line.add_run(f.suggestion)

    # 四、通过项简表
    doc.add_heading("四、通过项一览", level=1)
    passed = [f for f in findings if f.status == "pass"]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "规则 ID"
    table.rows[0].cells[1].text = "规则名称"
    for f in passed:
        cells = table.add_row().cells
        cells[0].text = f.rule_id
        cells[1].text = f.rule_name

    # 五、签署栏（自动质控 + 人工复核双签，0 外链纯文本）
    doc.add_heading("五、签署栏", level=1)
    doc.add_paragraph("自动质控：mrqc 病历内涵质控智能审核系统（规则引擎，数值结论全部来自确定性规则）。")
    doc.add_paragraph("质控医师（签字）：＿＿＿＿＿＿＿＿＿＿　　日期：＿＿＿＿年＿＿月＿＿日")
    doc.add_paragraph("复核医师（签字）：＿＿＿＿＿＿＿＿＿＿　　日期：＿＿＿＿年＿＿月＿＿日")
    note = doc.add_paragraph()
    note_run = note.add_run("说明：本报告为自动质控结果，须由质控医师复核签字后归档；"
                            "「待人工确认」结论为规则无法硬判、需人工核对原文的事项。")
    note_run.font.size = Pt(9)

    doc.save(str(out_path))
    return out_path
