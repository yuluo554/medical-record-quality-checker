"""docx 报告导出测试：免责声明/签署栏/依据关联内容、0 外链（rels 无 External）。"""

import zipfile
from pathlib import Path

import pytest

from mrqc.models import Evidence, FieldValue, PatientInfo, RecordCard
from mrqc.rules.finding import Basis, Finding, Severity

pytest.importorskip("docx", reason="report extra 未安装：pip install 'mrqc[report]'")

from mrqc.pipeline.stages import summarize  # noqa: E402
from mrqc.report import export_report  # noqa: E402


def _finding(rule_id="F-SIGN-01", status="fail", severity=Severity.HIGH,
             basis=True, suggestion="请补签名"):
    return Finding(
        rule_id=rule_id, rule_name="测试规则（出院记录未签名）", severity=severity,
        basis=Basis(document="病历书写基本规范", document_no="卫医政发〔2010〕11号",
                    clause="第三条", quote="病历书写应当客观、真实、准确。",
                    status="已核对") if basis else None,
        evidence=[Evidence(part="出院记录", quote="患者生命体征平稳，予以出院。")],
        suggestion=suggestion, status=status)


def _card() -> RecordCard:
    return RecordCard(patient=PatientInfo(
        name=FieldValue(value="张三（合成）"),
        department=FieldValue(value="普通外科"),
        age=FieldValue(value=45, unit="岁"),
    ))


def _doc_text(doc) -> str:
    """段落 + 表格单元格全文本（患者基本信息/汇总在表格里）。"""
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_export_report_minimal_record(tmp_path):
    """空参数卡 + 全 pass 结论也能出报告（免责声明/签署栏/通过项不缺席）。"""
    out = tmp_path / "r.docx"
    export_report(_card(), [], out, summary=summarize([], n_rules=33))
    assert out.is_file() and out.stat().st_size > 0
    import docx

    texts = _doc_text(docx.Document(str(out)))
    assert "病历内涵质控报告" in texts
    assert "不是临床决策支持工具" in texts          # 首页免责声明
    assert "全部规则通过" in texts
    assert "质控医师（签字）" in texts and "复核医师（签字）" in texts  # 签署栏


def test_export_report_findings_with_basis(tmp_path):
    """非 pass 结论逐条含分级/文号+条款+条文摘录/病历证据/整改建议。"""
    findings = [_finding(), _finding(rule_id="F-TIME-03", status="need_confirm",
                                     severity=Severity.NEED_CONFIRM, suggestion="")]
    out = tmp_path / "sub" / "r.docx"  # 父目录不存在也应自动创建
    export_report(_card(), findings, out, summary=summarize(findings, n_rules=33))
    import docx

    texts = _doc_text(docx.Document(str(out)))
    for expect in ("F-SIGN-01", "F-TIME-03", "卫医政发〔2010〕11号", "第三条",
                   "病历书写应当客观、真实、准确", "[出院记录]",
                   "较大风险", "待人工确认", "请补签名", "张三（合成）"):
        assert expect in texts, "报告缺少内容：%s" % expect


def test_export_report_zero_external_links(tmp_path):
    """0 外链纪律：docx 包内全部 rels 不得含 External 关系（无超链接/外链资源）。"""
    out = tmp_path / "r.docx"
    export_report(_card(), [_finding()], out, summary=summarize([_finding()]))
    with zipfile.ZipFile(str(out)) as z:
        for name in z.namelist():
            if name.endswith(".rels"):
                content = z.read(name).decode("utf-8")
                assert 'TargetMode="External"' not in content, \
                    "docx 含外链关系：%s" % name
